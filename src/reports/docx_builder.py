# src/reports/docx_builder.py
"""
Generador de informes DOCX (Word) — CN Hockey Performance Hub.

Contraparte editable de pdf_builder.py: mismo contenido, mismas secciones
(literalmente los mismos dataclasses SeccionFigura/SeccionTabla/SeccionFotos/
SeccionAsistente/SeccionAnalisis/SeccionTexto, importados de pdf_builder.py
— no se redefinen acá), pero en un .docx que Tomi puede editar a mano
después de generado (a pedido suyo: "pasamelo tanto en PDF como en Word por
si quiero hacer alguna modificación").

Decisión de arquitectura: se usa **python-docx**, no la librería docx-js/
Node.js que recomienda la skill de docx del entorno. Motivo: todo el resto
del stack de este repo (loaders, metrics, reports, la propia app Streamlit)
es Python — Tomi no conoce JS. Mantener la generación de reportes en un solo
lenguaje es más fácil de mantener para él a futuro que agregar una
dependencia de Node.js a un repo que hoy no la tiene. La calidad del
resultado es equivalente para lo que necesita este informe (texto, tablas,
imágenes de gráficos) — python-docx no tiene gotchas de docx-js como el
tamaño de página en DXA o ImageRun.type porque no es la misma librería.

Los gráficos (SeccionFigura) se recolorean para fondo blanco con la MISMA
función que ya usa pdf_builder.py (_recolorear_para_impresion) — se
reimporta desde ahí en vez de reimplementarla, para que un cambio de
paleta de impresión futuro no tenga que hacerse en dos lugares.
"""
import io
import re
from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from settings import ZSCORE_ALERTA
from src.reports.pdf_builder import (
    ANCHO_UTIL_CM, SeccionAnalisis, SeccionAsistente, SeccionFigura, SeccionFotos,
    SeccionTabla, SeccionTexto, _foto_cuadrada_buffer, _recolorear_para_impresion,
)

NAVY = RGBColor(0x0F, 0x2B, 0x5B)
TEXTO_GRIS = RGBColor(0x4B, 0x55, 0x63)
TEXTO_OSCURO = RGBColor(0x1F, 0x29, 0x37)
BLANCO = RGBColor(0xFF, 0xFF, 0xFF)

HEX_NAVY = "0F2B5B"
HEX_GRILLA_BG = "F3F4F6"

# Misma paleta clara que _ESTADO_PDF_CFG de pdf_builder.py — repetida acá
# (no importada) porque son solo 4 pares de strings hex, más simple que
# exponerla como público desde un módulo que hoy la tiene privada con "_".
_ESTADO_DOCX_CFG = {
    "Por debajo": {"texto": "0369A1", "bg": "E0F2FE"},
    "En rango":   {"texto": "15803D", "bg": "DCFCE7"},
    "Por encima": {"texto": "B91C1C", "bg": "FEE2E2"},
    "Sin dato":   {"texto": "6B7280", "bg": "F3F4F6"},
}

_TAG_BOLD = re.compile(r"<b>(.*?)</b>")


def _sombrear_celda(cell, hex_color: str) -> None:
    """Pinta el fondo de una celda de tabla — python-docx no expone esto en
    su API pública, hay que tocar el XML directo (mismo patrón que usa toda
    la comunidad de python-docx para shading de celdas)."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color.lstrip("#"))
    tcPr.append(shd)


def _run(paragraph, texto: str, bold: bool = False, color: RGBColor | None = None,
         size: float | None = None, italic: bool = False):
    if not texto:
        return None
    r = paragraph.add_run(texto)
    r.bold = bold
    r.italic = italic
    if color is not None:
        r.font.color.rgb = color
    if size is not None:
        r.font.size = Pt(size)
    return r


def _agregar_texto_con_negritas(paragraph, texto: str, color: RGBColor | None = None,
                                 size: float | None = None) -> None:
    """
    Los textos generados por reporte_semanal.py (resumen ejecutivo, FODA de
    periodización) traen marcado mínimo `<b>...</b>` — mismo criterio que ya
    interpreta reportlab.Paragraph en pdf_builder.py. python-docx no
    entiende HTML, así que acá se parsea a mano y se arman runs bold/no-bold
    alternados en el mismo párrafo.
    """
    pos = 0
    for m in _TAG_BOLD.finditer(texto):
        if m.start() > pos:
            _run(paragraph, texto[pos:m.start()], bold=False, color=color, size=size)
        _run(paragraph, m.group(1), bold=True, color=color, size=size)
        pos = m.end()
    if pos < len(texto):
        _run(paragraph, texto[pos:], bold=False, color=color, size=size)


def _quitar_tags(texto: str) -> str:
    return _TAG_BOLD.sub(r"\1", texto)


def _agregar_encabezado(doc: Document, titulo: str, subtitulo: str) -> None:
    p_titulo = doc.add_paragraph()
    _run(p_titulo, titulo, bold=True, color=NAVY, size=20)

    if subtitulo:
        p_sub = doc.add_paragraph()
        r = _run(p_sub, subtitulo, color=TEXTO_GRIS, size=10)
        if r is not None:
            r.italic = True

    p_fecha = doc.add_paragraph()
    from datetime import datetime
    _run(p_fecha, datetime.now().strftime("Generado el %d/%m/%Y a las %H:%M"), color=TEXTO_GRIS, size=8)
    doc.add_paragraph()


def _agregar_kpis(doc: Document, kpis: list[tuple[str, str]]) -> None:
    tabla = doc.add_table(rows=2, cols=len(kpis))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    ancho_col = Cm(ANCHO_UTIL_CM / len(kpis))
    for i, (label, valor) in enumerate(kpis):
        for fila in (0, 1):
            tabla.cell(fila, i).width = ancho_col
            _sombrear_celda(tabla.cell(fila, i), HEX_GRILLA_BG)

        p_label = tabla.cell(0, i).paragraphs[0]
        p_label.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p_label, label.upper(), color=TEXTO_GRIS, size=7)

        p_valor = tabla.cell(1, i).paragraphs[0]
        p_valor.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p_valor, valor, bold=True, color=NAVY, size=13)
    doc.add_paragraph()


def _agregar_titulo_seccion(doc: Document, titulo: str) -> None:
    p = doc.add_paragraph()
    _run(p, titulo, bold=True, color=NAVY, size=13)


def _agregar_tabla_df(doc: Document, df: pd.DataFrame) -> None:
    if df.empty:
        doc.add_paragraph("Sin datos para mostrar.")
        return
    tabla = doc.add_table(rows=1, cols=len(df.columns))
    tabla.style = "Table Grid"
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER

    for i, col in enumerate(df.columns):
        celda = tabla.rows[0].cells[i]
        _sombrear_celda(celda, HEX_NAVY)
        p = celda.paragraphs[0]
        _run(p, str(col), bold=True, color=BLANCO, size=9)

    for idx, (_, fila) in enumerate(df.astype(str).iterrows()):
        celdas = tabla.add_row().cells
        for i, valor in enumerate(fila):
            p = celdas[i].paragraphs[0]
            # Las tablas de reporte_semanal.py (FODA, plan de acción) pueden
            # traer <b>...</b> en la celda de "Hallazgo" — mismo tratamiento
            # que el resto del texto del informe.
            _agregar_texto_con_negritas(p, str(valor), size=9)
            if idx % 2 == 1:
                _sombrear_celda(celdas[i], HEX_GRILLA_BG)
    doc.add_paragraph()


def _agregar_tabla_asistente(doc: Document, df_evaluacion: pd.DataFrame, etiqueta_header: str) -> None:
    if df_evaluacion.empty:
        doc.add_paragraph("Sin datos para mostrar.")
        return
    metricas = list(dict.fromkeys(df_evaluacion["metrica"]))
    tabla = doc.add_table(rows=1, cols=len(metricas) + 1)
    tabla.style = "Table Grid"

    encabezados = [etiqueta_header] + [str(m) for m in metricas]
    for i, texto in enumerate(encabezados):
        celda = tabla.rows[0].cells[i]
        _sombrear_celda(celda, HEX_NAVY)
        _run(celda.paragraphs[0], texto, bold=True, color=BLANCO, size=9)

    for etiqueta, grupo in df_evaluacion.groupby("etiqueta", sort=False):
        fila_por_metrica = grupo.set_index("metrica")
        celdas = tabla.add_row().cells
        _run(celdas[0].paragraphs[0], str(etiqueta), bold=True, size=9)
        for col_idx, m in enumerate(metricas, start=1):
            if m not in fila_por_metrica.index:
                _run(celdas[col_idx].paragraphs[0], "—", size=9)
                continue
            r = fila_por_metrica.loc[m]
            if isinstance(r, pd.DataFrame):
                r = r.iloc[0]
            cfg = _ESTADO_DOCX_CFG.get(r["estado"], _ESTADO_DOCX_CFG["Sin dato"])
            valor_str = f"{r['valor_real']:.0f}" if pd.notna(r["valor_real"]) else "—"
            _run(celdas[col_idx].paragraphs[0], valor_str, bold=True,
                 color=RGBColor.from_string(cfg["texto"]), size=9)
            _sombrear_celda(celdas[col_idx], cfg["bg"])
    doc.add_paragraph()


def _agregar_analisis(doc: Document, titulo: str, analisis: dict) -> None:
    _agregar_titulo_seccion(doc, titulo)
    if not analisis["fortalezas"] and not analisis["debilidades"]:
        doc.add_paragraph("Sin fortalezas ni debilidades detectadas.")
        return

    if analisis["fortalezas"]:
        cfg_ok = _ESTADO_DOCX_CFG["En rango"]
        atipicas_por_nombre = {
            f["nombre"]: f["metricas_atipicas"] for f in analisis.get("fortalezas_atipicas", [])
        }
        p_label = doc.add_paragraph()
        _run(p_label, "Fortalezas", bold=True, size=10)
        _run(p_label, " — en rango en todas las métricas evaluadas:", size=10)

        p_chips = doc.add_paragraph()
        for i, nombre in enumerate(analisis["fortalezas"]):
            metricas = atipicas_por_nombre.get(nombre)
            etiqueta = nombre
            if metricas:
                detalle = ", ".join(f"{m['metrica']} z={m['z_score']:.1f}" for m in metricas)
                etiqueta = f"{nombre} ({detalle})"
            if i > 0:
                _run(p_chips, "  •  ", color=TEXTO_GRIS, size=9)
            _run(p_chips, etiqueta, bold=True, color=RGBColor.from_string(cfg_ok["texto"]), size=9)
        doc.add_paragraph()

    if analisis["debilidades"]:
        p_label = doc.add_paragraph()
        _run(p_label, "A vigilar", bold=True, size=10)
        for d in analisis["debilidades"]:
            p_nombre = doc.add_paragraph()
            _run(p_nombre, d["nombre"], bold=True, color=NAVY, size=11)
            _run(p_nombre, f"  ·  {d['posicion']}", color=TEXTO_GRIS, size=8)

            p_metricas = doc.add_paragraph()
            todas_metricas = d["metricas_fuera"] + d["metricas_en_rango"]
            for i, m in enumerate(todas_metricas):
                cfg = _ESTADO_DOCX_CFG.get(m["estado"], _ESTADO_DOCX_CFG["Sin dato"])
                valor = f"{m['valor_real']:.0f}" if pd.notna(m["valor_real"]) else "—"
                z = m.get("z_score")
                marca = ""
                if pd.notna(z):
                    marca = f" (z={z:.1f})" if abs(z) >= ZSCORE_ALERTA else ""
                if i > 0:
                    _run(p_metricas, "   ", size=8)
                _run(p_metricas, f"{m['metrica']}: {valor}{marca}", bold=True,
                     color=RGBColor.from_string(cfg["texto"]), size=8)

            for reco in d["recomendaciones"]:
                p_reco = doc.add_paragraph()
                _run(p_reco, reco, color=TEXTO_OSCURO, size=8)
            doc.add_paragraph()


def _agregar_figura(doc: Document, figura, alto_cm: float) -> None:
    fig_clara = _recolorear_para_impresion(figura)
    ancho_px = 1400
    alto_px = int(ancho_px * (alto_cm / ANCHO_UTIL_CM))
    png_bytes = fig_clara.to_image(format="png", width=ancho_px, height=alto_px, scale=2)
    doc.add_picture(io.BytesIO(png_bytes), width=Cm(ANCHO_UTIL_CM))
    doc.add_paragraph()


def _agregar_fotos(doc: Document, fotos: list[tuple[str, Path | str | None]], lado_cm: float) -> None:
    tabla = doc.add_table(rows=2, cols=len(fotos))
    tabla.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, (nombre, ruta) in enumerate(fotos):
        buffer = _foto_cuadrada_buffer(Path(ruta)) if ruta and Path(ruta).exists() else None
        p_img = tabla.cell(0, i).paragraphs[0]
        p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if buffer:
            run = p_img.add_run()
            run.add_picture(buffer, width=Cm(lado_cm))
        else:
            _run(p_img, "(sin foto)", size=8)
        p_nombre = tabla.cell(1, i).paragraphs[0]
        p_nombre.alignment = WD_ALIGN_PARAGRAPH.CENTER
        _run(p_nombre, nombre, bold=True, size=9)
    doc.add_paragraph()


def generar_docx_reporte(
    titulo: str,
    subtitulo: str,
    secciones: list[SeccionFigura | SeccionTabla | SeccionFotos | SeccionAsistente
                    | SeccionAnalisis | SeccionTexto],
    kpis: list[tuple[str, str]] | None = None,
) -> bytes:
    """
    Arma el informe DOCX y devuelve sus bytes — misma firma que
    generar_pdf_reporte() de pdf_builder.py, para que generar_reporte_semanal.py
    llame a las dos con la MISMA lista de secciones y no tenga que armar el
    contenido dos veces.
    """
    doc = Document()
    for section in doc.sections:
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)
        section.top_margin = Cm(1.8)
        section.bottom_margin = Cm(1.6)

    estilo_normal = doc.styles["Normal"]
    estilo_normal.font.name = "Calibri"
    estilo_normal.font.size = Pt(10)

    _agregar_encabezado(doc, titulo, subtitulo)
    if kpis:
        _agregar_kpis(doc, kpis)

    for seccion in secciones:
        if isinstance(seccion, SeccionAnalisis):
            _agregar_analisis(doc, seccion.titulo, seccion.analisis)
            continue

        _agregar_titulo_seccion(doc, seccion.titulo)
        if isinstance(seccion, SeccionFigura):
            _agregar_figura(doc, seccion.figura, seccion.alto_cm)
        elif isinstance(seccion, SeccionAsistente):
            _agregar_tabla_asistente(doc, seccion.df_evaluacion, seccion.etiqueta_header)
        elif isinstance(seccion, SeccionTabla):
            _agregar_tabla_df(doc, seccion.df)
        elif isinstance(seccion, SeccionFotos):
            _agregar_fotos(doc, seccion.fotos, seccion.lado_cm)
        elif isinstance(seccion, SeccionTexto):
            if not seccion.parrafos:
                doc.add_paragraph("Sin datos para mostrar.")
            else:
                for parrafo in seccion.parrafos:
                    p = doc.add_paragraph()
                    _agregar_texto_con_negritas(p, parrafo, color=TEXTO_OSCURO, size=10.5)
        doc.add_paragraph()

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()
